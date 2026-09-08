"""Turkish export parity, value integrity, and language-independent color checks."""
from copy import deepcopy
from io import BytesIO
import re

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pytest

from cerai_i18n import translate_text
import reports
from test_step10_canonical_reports import _payload


def _word_text(document):
    return '\n'.join([p.text for p in document.paragraphs] + [
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    ])


def _pdf_text(data):
    return re.sub(r'\s+', ' ', ' '.join(p.extract_text() for p in PdfReader(BytesIO(data)).pages))


def test_turkish_exports_translate_current_report_sections_and_preserve_snapshot():
    payload = _payload(ml7_bad_k1_d=40, ml7_bad_k2_d=45)
    payload['locale'] = 'tr'
    payload['patient'].update(name='No Source Şule Işık', id='PASS-NO-035', reviewer='Dr. Çağrı Şen')
    original = deepcopy(payload)
    document = Document(BytesIO(reports.build_docx(payload)))
    word, pdf = _word_text(document), _pdf_text(reports.build_pdf(payload))
    for content in (word, pdf):
        for required in (
            'Genel karar', 'OD — UYGUN', 'Cerrahi güvenlik', 'Cerrahi planlama',
            'Kararın gerekçesi', 'Pentacam değerleri ve kaynak bilgileri',
                'Topografi kategorisi', 'Astigmatik uyumsuzluk doğrulaması', 'Gözler arası puan 0/5.',
            'Kornea kalınlık haritası morfolojisi: değerlendirilmedi; cerrah değerlendirmesi gerekir;',
            'Göreli kalınlık haritası: değerlendirilmedi; cerrah değerlendirmesi gerekir;',
            'PTI/CTSP kalınlık profili morfolojisi: değerlendirilmedi; cerrah değerlendirmesi gerekir;',
            'Mitomisin-C önerisi', 'NORMAL / UYGUN', 'yalnızca bilgilendirme',
            'UYGUN — CER-AI nihai birleştirme ölçütleri karşılandı',
            'Kaynak: Axial/Sagittal Curvature (Front).',
            'No Source Şule Işık', 'PASS-NO-035', 'Dr. Çağrı Şen',
            'SHOW_2_CORNEA_BACK / show2.png', 'BAD_D_STRIP / OD-bad.png',
                'CER-AI-2026-09-08-PS3-DISPARITY-SEPARATED-ML7-HINGE-V3',
        ):
            assert required in content
        for obsolete in ('Canonical result', 'Procedure disposition', 'selected_plan',
                         'information only', 'Recommendation only;', 'Inter-eye score'):
            assert obsolete not in content
    assert document.core_properties.language == 'tr-TR'
    assert payload == original


@pytest.mark.parametrize('status,label,color', [
    ('PASS', 'UYGUN', reports.GREEN_FILL),
    ('PASS WITH CAUTION', 'DİKKATLE UYGUN', reports.AMBER_FILL),
    ('CAUTION', 'DİKKAT', reports.AMBER_FILL),
    ('STOP-DEFER', 'DURDUR-ERTELE', reports.RED_FILL),
])
def test_translated_status_keeps_canonical_colors_in_both_formats(status,label,color):
    # Presentation-only fixture: no risk is inferred from the translated label.
    payload = _payload()
    payload['locale'] = 'tr'
    payload['decision']['status'] = status
    for eye in payload['decision']['eyes']:
        eye['report_payload']['status'] = status
        eye['report_payload']['randleman']['status'] = status
    doc = Document(BytesIO(reports.build_docx(payload)))
    cell = doc.tables[0].rows[2].cells[3]
    assert cell.text == label
    assert cell._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == color
    erss_cell = doc.tables[1].rows[-1].cells[1]
    assert erss_cell.text == label
    assert erss_cell._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == color
    reader = PdfReader(BytesIO(reports.build_pdf(payload)))
    assert f'OD — {label}' in _pdf_text(reports.build_pdf(payload))
    from reportlab.lib.colors import HexColor
    rgb = HexColor('#'+color)
    # The first page must paint the canonical disposition's fill color.
    from pypdf.generic import ContentStream
    paints = [args for args,op in ContentStream(reader.pages[0].get_contents(),reader).operations if op == b'rg']
    assert any(all(abs(float(a)-b)<0.00001 for a,b in zip(args,(rgb.red,rgb.green,rgb.blue))) for args in paints)


def test_dynamic_findings_keep_values_comparisons_and_units():
    messages = [
        'Manifest astigmatism 2.75 D and topographic astigmatism 2 D are both <=3.00 D; comparison inactive, no PS3 risk factor.',
        'Front-map SRAX 20.1° > 20°. Source: Axial/Sagittal Curvature (Front).',
        'Astigmatism difference 1.25 D; BAD flat-axis versus minus-cylinder manifest axis difference 10.1°.',
        'Anterior Km 48.5 D is 48-50 D.',
        'F.Ele.Th 5 µm <=12 and B.Ele.Th 10 µm <=15.',
    ]
    token = r'[<>]=?|[+-]?\d+(?:\.\d+)?|µm|°'
    for message in messages:
        translated = translate_text(message, 'tr')
        assert translated != message
        assert re.findall(token,message) == re.findall(token,translated)


def test_joined_planning_notes_and_nested_drivers_are_translated():
    note = ('Recommendation only; surgeon must verify anatomy, device setup, and the active ML7 manual before use.; '
            'If the selected ring leaks, does not hold vacuum, or is too large, the active ML7 reference directs selection of one smaller ring.')
    translated = translate_text(note,'tr')
    assert 'Yalnızca öneridir' in translated and 'bir küçük halka' in translated
    assert 'Recommendation' not in translated and 'If the selected' not in translated
    driver = 'key: procedural_safety; status: STOP-DEFER; detail: Independent tissue/refractive safety gates'
    assert translate_text(driver,'tr') == 'Ölçüt: Cerrahi güvenlik; Durum: DURDUR-ERTELE; Açıklama: Bağımsız doku ve refraktif güvenlik koşulları'
    assert translate_text('OD: LASIK failed. Now evaluating PRK.','tr') == 'OD: LASIK uygun bulunmadı. PRK değerlendiriliyor.'
    assert translate_text('PRK_PTA_percent','tr') == 'PRK PTA (%)'


def test_english_translation_is_identity_and_protected_cells_remain_literal():
    for original in ('selected_plan','No Source','PASS-NO.png','SHOW_2_CORNEA_BACK / No.png'):
        assert translate_text(original,'en') == original
        assert reports._cell_text(original,'tr',literal=True) == original
    assert reports._cell_text('Not documented','tr',literal=True) == 'Belgelenmedi'


def test_turkish_bad_classifications_and_selected_plan_keep_colors():
    payload = _payload(Df=2.7, Db=1.8)
    payload['locale'] = 'tr'
    doc = Document(BytesIO(reports.build_docx(payload)))
    bad = next(t for t in doc.tables if t.cell(0,0).text == 'Parametre' and t.cell(0,2).text == 'Yorum')
    df = next(row for row in bad.rows if row.cells[0].text == 'Df')
    db = next(row for row in bad.rows if row.cells[0].text == 'Db')
    assert df.cells[2].text == 'ANORMAL / >= 2.60; yalnızca bilgilendirme'
    assert db.cells[2].text == 'ŞÜPHELİ / 1.60 ile < 2.60; yalnızca bilgilendirme'
    assert df.cells[1]._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == reports.RED_FILL
    assert db.cells[1]._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == reports.AMBER_FILL
    selected = next(row for t in doc.tables for row in t.rows if row.cells[0].text == 'Seçilen plan')
    assert selected.cells[1].text == 'Plan A'
    assert selected.cells[1]._tc.get_or_add_tcPr().find(qn('w:shd')).get(qn('w:fill')) == reports.GREEN_FILL
    assert 'pachymetry' not in translate_text('Inter-eye score 1/5; exceeded: thinnest pachymetry.', 'tr')
